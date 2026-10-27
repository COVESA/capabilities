"""Generate the editable COVESA talk. Requires python-pptx in the local environment."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


OUTPUT = Path(__file__).resolve().parent
SLIDES = [
    {
        "title": "Vehicle Capabilities",
        "subtitle": "A runnable power-window reference example",
        "section": "COVESA All Member Meeting",
        "seconds": 45,
        "kind": "title",
        "items": ["Semantic intent", "Stable API boundaries", "Executable models"],
        "notes": "Introduce the talk as a reference example and a discussion proposal, not an adopted COVESA standard. We will follow one window command from an application to a simulated motor. The goal is to make interfaces and behavior concrete enough for members to challenge and improve. Add the approved presenter name and meeting date before presenting.",
    },
    {
        "title": "Applications Need Intent, Not Wiring",
        "subtitle": "Move the window to 75% open, without knowing the motor driver.",
        "section": "The Problem",
        "seconds": 75,
        "items": ["Signals describe values", "Transport APIs describe access", "Capabilities describe behavior"],
        "notes": "A window is deliberately ordinary. An application should not need a bus identifier, supplier pinout, or ECU address to request a position. A transport API alone does not settle ownership, inhibition, completion, or failure semantics. A capability is the place to give that request a domain meaning. This complements signal-level interoperability rather than replacing it.",
    },
    {
        "title": "A Small, Deliberate Reference Scope",
        "subtitle": "One implemented capability. Enough surrounding structure to run it.",
        "section": "Scope",
        "seconds": 90,
        "items": ["Power windows: four semantic zones", "HTTP adapter, catalog and explorer", "Simulator binding and motor model"],
        "callout": "Vehicle state and power management are platform stand-ins, not production services.",
        "notes": "The runnable scope is windows only. Other domains in the architectural picture are context, not released capability implementations. The composition root wires the bus, vehicle-state stub, simulator, manager, and window orchestrator. A power stub is present but the manager does not request power authorization. Do not imply that this example implements the full vehicle architecture.",
    },
    {
        "title": "Two Boundaries, One Capability",
        "subtitle": "Separate caller intent from deployment-specific actuation.",
        "section": "Architecture",
        "seconds": 120,
        "kind": "architecture",
        "items": ["Application / HMI", "Northbound API", "Window capability", "Southbound port", "Simulator + motor model"],
        "notes": "Read the diagram top to bottom. The northbound API accepts semantic commands and returns typed state. The manager owns domain gates and dispatch. The southbound port exposes abstract actuation and position read-back. The simulator is the implemented binding; a hardware binding is a future integration point. Deployment addressing remains below the port. State events provide feedback to callers. These boundaries are architectural contracts, not necessarily process or network boundaries.",
    },
    {
        "title": "The Northbound Contract",
        "subtitle": "What callers ask for, and what they can reason about.",
        "section": "Northbound",
        "seconds": 90,
        "items": ["Move: zone + target + request identity", "Preflight and explicit release", "State: position, status and typed reason"],
        "callout": "API convention: 0% = closed; 100% = open.",
        "notes": "Describe MoveWindowCommand and WindowState. The adapter derives the request source from its configured role and rejects a source field in the payload. That is a demonstration configuration, not a credential verification system. Runtime endpoints are under /vehicle/window. The checked-in OpenAPI paths and response codes currently differ from the adapter and need alignment before consumers generate clients. Avoid presenting that drift as a standardized contract.",
    },
    {
        "title": "A Request Has Five Phases",
        "subtitle": "Make execution order visible and inspectable.",
        "section": "Behavior",
        "seconds": 90,
        "kind": "pipeline",
        "items": ["Arbitrate", "Precondition", "Safety checks", "Actuate", "Publish"],
        "notes": "The manager executes arbitration, preconditions, safety checks, actuator dispatch, and state publication. Current gates check power OFF, OTA activity, undervoltage, and target range. Timeout and explicit release use cancellation and position read-back support. The term safety checks is not an ISO 26262 safety case or certification claim. Priority logic exists, but the manager holds one lock throughout a move, so immediate priority preemption and independent zone concurrency should not be claimed as proven. The explorer's phase replay is an explanatory animation, not live execution tracing.",
    },
    {
        "title": "The Southbound Contract",
        "subtitle": "A binding changes; the caller's vocabulary does not.",
        "section": "Southbound",
        "seconds": 90,
        "items": ["move_window(zone, target_pct)", "get_window_position(zone)", "describe_actuator(capability)"],
        "callout": "Timeout or release requires a binding that can stop safely on cancellation.",
        "notes": "The WindowActuatorPort is an abstract interface. The simulator maps API percentages to model millimeters and reports the resulting position. A deployment profile describes electrical, travel, load, and geometry quantities without hardware addresses. A real binding would implement transport and feedback. The simulator runs physics synchronously inside its async method, so it is not evidence of interruptible hardware timing. No real ECU integration is demonstrated here.",
    },
    {
        "title": "An Executable Plant Model",
        "subtitle": "Connect interface intent to physical behavior.",
        "section": "Model",
        "seconds": 90,
        "kind": "pipeline",
        "items": ["DC motor", "Actuator", "Linkage", "Window load"],
        "callout": "Simulated KPIs: move latency, current, force, travel and energy.",
        "notes": "Explain the four-stage bond-graph representation: electrical drive and motor, rotary-to-linear actuation, linkage, and window load. The model can generate a trajectory, compute KPIs, and respond to an imposed obstruction in model-level tests. Model position is the opposite of API openness: 0 mm is open and 300 mm is closed for the default model. Parameters and outputs are illustrative simulation data, not measured supplier specifications or vehicle validation evidence. No production anti-pinch timing guarantee is established.",
    },
    {
        "title": "Capabilities Complement VSS",
        "subtitle": "Keep signal references, behavior and deployment separate.",
        "section": "COVESA Alignment",
        "seconds": 90,
        "items": ["Signal reference: VSS position path", "Behavior: MoveWindowCommand / WindowState", "Deployment: actuator binding and profile"],
        "callout": "Vehicle.Cabin.Door.Row1.DriverSide.Window.Position",
        "notes": "VSS supplies a shared signal vocabulary. The example attaches VSS paths to capability metadata and relevant operations, while defining command and result semantics separately. This is VSS-referenced metadata, not a demonstrated VSS server integration or a conformance certification. Agreeing units, ranges, and direction conventions is essential; review the exact mapping against the chosen VSS release. The catalog can help discover interfaces and model parameters without embedding transport addresses.",
    },
    {
        "title": "Demo: Follow One Window Command",
        "subtitle": "Three minutes: discover, explain, execute, inspect.",
        "section": "Live Demo",
        "seconds": 180,
        "items": ["Catalog: select Window Movement Control", "Boundary: step through the phases", "Explorer: move to 75% and inspect state"],
        "callout": "Local demo: http://127.0.0.1:8770/vehicle/catalog/",
        "notes": "0:00-0:30: Open http://127.0.0.1:8770/vehicle/catalog/?cap=window&tab=overview and show the descriptor and ports. 0:30-1:15: Open ?cap=window&tab=boundary, choose a move operation, and step through the explanatory replay. State explicitly that replay is not live instrumentation. 1:15-2:00: Open ?main_tab=explorer, choose DRIVER_FRONT and 75%, execute, then show the response and event log. The configured adapter role controls source; do not rely on a selector to switch identity. 2:00-3:00: Open /vehicle/api-explorer/ and show the physics playground trajectory and KPIs. Explain that playground simulation is a separate model run, not motor telemetry. If the demo is unavailable, use the deck's architecture and pipeline diagrams and the following HTTP fallback: POST /vehicle/window/move with {\"zone\":\"DRIVER_FRONT\",\"target_pct\":75}; GET /vehicle/window/state?zone=DRIVER_FRONT. Rehearse on the actual machine first. Start using python -m capability_framework.app --no-browser.",
    },
    {
        "title": "What The Example Demonstrates",
        "subtitle": "Concrete evidence, without overstating readiness.",
        "section": "Evidence",
        "seconds": 90,
        "items": ["Runnable HTTP-to-simulator path", "Tests for model, API and explorer", "Introspectable interfaces and model"],
        "callout": "Not demonstrated: real hardware, certified safety or production authentication.",
        "notes": "The repository has tests for target movement, model KPIs, obstruction detection, power-off inhibition, HTTP operations, catalog rendering, diagrams, and playground simulation. These establish behavior within a reference example, not production readiness. Do not cite a test count without running the checked-out revision. The window orchestrator issues manager commands sequentially even for its SYNC convenience mode. The runtime returns MOVING after its simulated call completes; completion and status semantics remain a collaboration question. Calibration profile types exist, but automatic outcome reporting is not wired in the current main revision.",
    },
    {
        "title": "Candidate Topics For Collaboration",
        "subtitle": "Agree behavior at the boundaries before prescribing deployment.",
        "section": "Discussion Proposal",
        "seconds": 60,
        "items": ["Command lifecycle and completion semantics", "Cancellation, ownership and typed failures", "Port contracts, units and VSS mappings"],
        "notes": "Present these as candidate topics for member discussion. No adopted standard or workgroup commitment is implied. Which elements should be portable across deployments? How should asynchronous completion and state observation work? What must a binding guarantee on cancellation? What verification examples should accompany a contract? Keep transport choices and hardware topology as extension concerns unless members identify a compelling interoperability need.",
    },
    {
        "title": "From Reference To Shared Artifact",
        "subtitle": "Small scope, explicit gaps, repeatable review.",
        "section": "Next Steps",
        "seconds": 60,
        "items": ["Align contract, catalog and runtime", "Add conformance and fault-injection cases", "Review provenance and publication approval"],
        "notes": "Propose a practical sequence. First align the OpenAPI paths, statuses, and schemas with runtime behavior. Then exercise arbitration, interruptibility, cancellation, fault reporting, and model limits. Validate a second binding before claiming hardware portability. Publication requires the owner's review of code and parameter provenance, copyright permissions, and licensing. The repository LICENSE is currently a placeholder while package metadata declares Apache-2.0; this must be reconciled. A keyword scan cannot establish absence of proprietary information. The deck avoids employee names, internal repositories, and internal rule identifiers, but remains subject to release review.",
    },
    {
        "title": "Which Boundary Should We Agree First?",
        "subtitle": "One window. Two interfaces. A concrete starting point.",
        "section": "Discussion",
        "seconds": 30,
        "kind": "title",
        "items": ["Caller semantics", "Binding guarantees", "Verification evidence"],
        "notes": "Close with the central invitation: use this small runnable example to decide which boundary semantics deserve shared agreement. Ask members to nominate one missing guarantee, one binding scenario, or one verification case. The planned talk is 20 minutes including the demo; any extended questions need additional agenda time. Do not announce a public repository link until publication is approved.",
    },
]

INK = "17212B"
MUTED = "586570"
PAPER = "F4F7F9"
TEAL = "007F82"
CORAL = "C44B38"


def color(value):
    return RGBColor.from_string(value)


def rectangle(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background()
    return shape


def text(slide, left, top, width, height, value, size=24, fill=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color(fill)
        paragraph.space_after = Pt(12)
    return box


def build():
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    deck.core_properties.title = "Vehicle Capabilities: A Runnable Power-Window Reference Example"
    deck.core_properties.subject = "COVESA All Member Meeting - 20-minute discussion presentation"
    deck.core_properties.author = ""
    deck.core_properties.keywords = "COVESA, VSS, capabilities, windows, reference example"
    notes = ["# COVESA All Member Meeting: Presenter Guide", "",
             "20 minutes total, including a 3-minute demo. Extended Q&A is additional.",
             "Draft for technical and publication review. Add approved presenter/date in PowerPoint.",
             "No proprietary status, safety certification, or COVESA endorsement is asserted.", ""]
    elapsed = 0
    for index, spec in enumerate(SLIDES, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color(PAPER)
        rectangle(slide, 0, 0, 0.16, 7.5, TEAL)
        text(slide, 0.65, 0.4, 11.8, 0.4, spec["section"].upper(), 13, TEAL, True)
        text(slide, 0.65, 1.05, 12, 1.0, spec["title"], 36, INK, True)
        text(slide, 0.65, 2.15, 11.8, 0.8, spec["subtitle"], 23, MUTED)
        kind = spec.get("kind", "list")
        if kind == "architecture":
            for item_index, item in enumerate(spec["items"]):
                top = 3.0 + item_index * 0.64
                rectangle(slide, 0.85, top, 7.5, 0.48, TEAL if item_index in (1, 3) else INK)
                text(slide, 1.05, top + 0.06, 7.1, 0.36, item, 18, "FFFFFF", True)
            text(slide, 9.0, 3.0, 3.5, 2.9,
                 "Intent flows down\n\nPosition + typed state return\n\nHardware addresses stay below the port", 22, TEAL)
        elif kind == "pipeline":
            count = len(spec["items"])
            width = 11.7 / count
            for item_index, item in enumerate(spec["items"]):
                left = 0.75 + item_index * width
                rectangle(slide, left, 3.3, width - 0.2, 1.5, INK)
                text(slide, left + 0.18, 3.45, width - 0.55, 0.4, f"{item_index + 1:02}", 17, "5ED4CD", True)
                text(slide, left + 0.18, 4.05, width - 0.55, 0.6, item, 21, "FFFFFF", True)
        else:
            for item_index, item in enumerate(spec["items"]):
                top = 3.2 + item_index * 0.85
                text(slide, 0.75, top, 0.65, 0.55, f"{item_index + 1:02}", 20, CORAL, True)
                text(slide, 1.65, top, 10.7, 0.7, item, 28, INK, kind == "title")
        if spec.get("callout"):
            rectangle(slide, 0.65, 6.15, 12, 0.58, "E3EEEF")
            text(slide, 0.85, 6.27, 11.6, 0.4, spec["callout"], 16, TEAL)
        text(slide, 0.65, 7.03, 10.6, 0.25, "COVESA All Member Meeting | Window capability reference example", 10, MUTED)
        text(slide, 11.7, 6.99, 0.9, 0.3, f"{index:02} / {len(SLIDES)}", 11, MUTED)
        start = f"{elapsed // 60:02}:{elapsed % 60:02}"
        elapsed += spec["seconds"]
        end = f"{elapsed // 60:02}:{elapsed % 60:02}"
        note = f"Timing: {start}-{end} ({spec['seconds']} seconds)\n\n{spec['notes']}"
        slide.notes_slide.notes_text_frame.text = note
        notes.extend([f"## {index}. {spec['title']} ({start}-{end})", "", spec["notes"], ""])
    assert elapsed == 1200
    deck.save(OUTPUT / "covesa-window-capabilities.pptx")
    (OUTPUT / "covesa-window-capabilities-notes.md").write_text("\n".join(notes), encoding="utf-8")
    print(f"Generated {len(SLIDES)} slides; scheduled duration: {elapsed // 60} minutes")


if __name__ == "__main__":
    build()